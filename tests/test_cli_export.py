from __future__ import annotations

import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from unit_economics.cli import main
from unit_economics.demo import DemoFiles


def report_args(demo_files: DemoFiles, out: Path, *extra: str) -> list[str]:
    return [
        "report", "--wb", *map(str, demo_files.wb), "--ozon", *map(str, demo_files.ozon),
        "--costs", str(demo_files.costs), "-o", str(out), *extra,
    ]  # fmt: skip


def test_demo_command(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["demo", "--out", str(tmp_path / "demo")]) == 0
    workbook = load_workbook(tmp_path / "demo" / "юнит-экономика.xlsx")
    assert workbook.sheetnames == ["Итоги", "По артикулам", "Общие расходы", "Проблемы"]
    assert [cell.value for cell in workbook["Итоги"][3]] == ["Показатель", "WB", "Ozon", "Итого"]
    assert workbook["По артикулам"].max_row == 11  # заголовок и 5 товаров × 2 площадки
    printed = capsys.readouterr().out
    assert "Чистая прибыль" in printed
    assert "Не задана себестоимость: SK-LN" in printed


def test_report_command_with_tax_and_ads(demo_files: DemoFiles, tmp_path: Path) -> None:
    out = tmp_path / "итог.xlsx"
    assert main(report_args(demo_files, out, "--tax", "usn-profit", "--rate", "15", "--ads-wb", "10 000")) == 0
    sheet = load_workbook(out)["Итоги"]
    rows = {row[0]: row[1:] for row in sheet.iter_rows(min_row=4, values_only=True)}
    assert rows["Реклама (введена вручную)"][0] == -10_000  # WB
    assert rows["Налог"][0] is None  # «доходы минус расходы»: только в «Итого»
    assert rows["Налог"][2] is not None
    articles = load_workbook(out)["По артикулам"]
    assert articles["G2"].number_format == '#,##0.00 "₽"'


def test_wrong_options_and_files(demo_files: DemoFiles, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit):
        main(["report", "--wb", str(demo_files.wb[0]), "--tax", "osno"])
    with pytest.raises(SystemExit):
        main(["report"])
    assert "хотя бы один отчёт" in capsys.readouterr().err
    broken = tmp_path / "wb.xlsx"
    broken.write_text("не Excel", encoding="utf-8")
    assert main(["report", "--wb", str(broken), "-o", str(tmp_path / "x.xlsx")]) == 2
    assert "Не удалось открыть" in capsys.readouterr().err


def test_report_open_in_excel(
    demo_files: DemoFiles, tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    out = tmp_path / "итог.xlsx"
    out.write_bytes(b"old")
    replace = os.replace

    def locked(source: str | Path, target: str | Path) -> None:
        if str(target).endswith(".xlsx"):
            raise PermissionError(13, "Отказано в доступе")
        replace(source, target)

    monkeypatch.setattr(os, "replace", locked)
    assert main(report_args(demo_files, out)) == 1
    assert "Закройте его" in capsys.readouterr().err
    assert out.read_bytes() == b"old"
