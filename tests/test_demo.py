"""Разбор демо-файлов сверяется с расчётом прямо по ячейкам — независимо от кода парсеров."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from unit_economics.demo import DemoFiles
from unit_economics.ozon import parse_ozon
from unit_economics.pnl import Settings, build_pnl
from unit_economics.wb import parse_wb

WB_EXPENSES = (
    "Услуги по доставке товара покупателю", "Общая сумма штрафов", "Хранение", "Удержания", "Платная приемка",
)  # fmt: skip


def wb_by_formula(path: Path) -> tuple[int, int, int]:
    """Сумма к оплате по формуле WB, проданные штуки и число строк продаж."""
    rows = load_workbook(path).worksheets[0].iter_rows(values_only=True)
    at = {name: index for index, name in enumerate(next(rows))}
    total, sold, sale_rows = Decimal(0), 0, 0
    for row in rows:

        def money(name: str, row: tuple[object, ...] = row) -> Decimal:
            return Decimal(str(row[at[name]] or 0))

        pay = money("К перечислению Продавцу за реализованный Товар")
        total += -pay if row[at["Тип документа"]] == "Возврат" else pay
        total += money("Доплаты") - sum((money(name) for name in WB_EXPENSES), Decimal(0))
        if row[at["Обоснование для оплаты"]] == "Продажа":
            sold += int(str(row[at["Кол-во"]]))
            sale_rows += 1
    return int(total * 100), sold, sale_rows


@pytest.mark.parametrize("week", [0, 1])
def test_wb_parser_matches_wb_payout_formula(demo_files: DemoFiles, week: int) -> None:
    path = demo_files.wb[week]
    report = parse_wb(path)
    payout, sold, sale_rows = wb_by_formula(path)
    assert sum(operation.payout for operation in report.operations) == payout
    assert sum(operation.sold for operation in report.operations) == sold
    assert sold > sale_rows  # в демо есть строки «Кол-во» = 2


def test_ozon_parser_matches_totals_row(demo_files: DemoFiles) -> None:
    path = demo_files.ozon[0]
    rows = list(load_workbook(path).worksheets[0].iter_rows(values_only=True))
    header = next(index for index, row in enumerate(rows) if "Тип начисления" in row)
    at = {name: index for index, name in enumerate(rows[header])}
    assert rows[-1][at["ID начисления"]] == "Итого"
    report = parse_ozon(path)
    assert sum(operation.payout for operation in report.operations) == int(Decimal(str(rows[-1][at["Итого"]])) * 100)


def test_demo_tells_its_story(demo_files: DemoFiles) -> None:
    reports = [parse_wb(path) for path in demo_files.wb] + [parse_ozon(path) for path in demo_files.ozon]
    from unit_economics.costs import load_costs

    result = build_pnl(reports, load_costs(demo_files.costs), Settings())
    towel = {row.platform: row for row in result.articles if row.article == "PT-VF"}
    assert towel["WB"].margin_profit > 0  # дешёвое полотенце выгодно на WB…
    assert towel["Ozon"].margin_profit < 0  # …и убыточно на Ozon
    assert result.missing_costs == ("SK-LN",)
    assert any("Коррекция продаж" in problem for problem in result.problems)
    assert any("Корректировка стоимости услуги" in problem for problem in result.problems)
