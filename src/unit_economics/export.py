"""Отчёт в Excel (.xlsx): итоги по площадкам, артикулы, общие расходы, проблемы.

Не устаревший .xls в заранее заданную папку: путь выбирает пользователь, файл пишется во временный и подменяет старый.
"""

from __future__ import annotations

import os
from collections.abc import Sequence
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from unit_economics.fmt import Kind, Value
from unit_economics.pnl import ARTICLE_COLUMNS, SUMMARY_ROWS, Pnl

FORMATS: dict[Kind, str] = {"money": '#,##0.00 "₽"', "count": "#,##0", "percent": "0.0%", "text": "General"}
TITLE_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2A5C8A")
LOSS_FILL = PatternFill("solid", fgColor="FBE3E3")
MISSING_FILL = PatternFill("solid", fgColor="FFF4D6")


class ExportError(Exception):
    pass


def excel_value(value: Value, kind: Kind) -> object:
    if value is None:
        return None
    if kind == "money":
        return Decimal(int(value)) / 100
    return value


def _header(sheet: Worksheet, row: int, titles: Sequence[str], widths: Sequence[float]) -> None:
    for column, (title, width) in enumerate(zip(titles, widths, strict=True), start=1):
        cell = sheet.cell(row, column, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width


def build_workbook(pnl: Pnl) -> Workbook:
    workbook = Workbook()
    del workbook["Sheet"]

    summary = workbook.create_sheet("Итоги")
    summary.cell(1, 1, f"Юнит-экономика — {pnl.title}").font = TITLE_FONT
    _header(summary, 3, ["Показатель", *(t.platform for t in pnl.totals)], [40, *([18] * len(pnl.totals))])
    for row, (title, getter, kind) in enumerate(SUMMARY_ROWS, start=4):
        summary.cell(row, 1, title)
        for column, totals in enumerate(pnl.totals, start=2):
            summary.cell(row, column, excel_value(getter(totals), kind)).number_format = FORMATS[kind]
    summary.freeze_panes = "B4"

    articles = workbook.create_sheet("По артикулам")
    widths = [10, 12, 34, *([13] * (len(ARTICLE_COLUMNS) - 3))]
    _header(articles, 1, [title for title, _, _ in ARTICLE_COLUMNS], widths)
    for row, item in enumerate(pnl.articles, start=2):
        for column, (title, value_of, kind) in enumerate(ARTICLE_COLUMNS, start=1):
            cell = articles.cell(row, column, excel_value(value_of(item), kind))
            cell.number_format = FORMATS[kind]
            if item.margin_profit < 0:
                cell.fill = LOSS_FILL
            elif item.unit_cost is None and title.startswith("Себестоимость"):
                cell.fill = MISSING_FILL
    articles.freeze_panes = "D2"
    if pnl.articles:
        articles.auto_filter.ref = f"A1:{get_column_letter(len(ARTICLE_COLUMNS))}{len(pnl.articles) + 1}"

    general = workbook.create_sheet("Общие расходы")
    _header(general, 1, ["Площадка", "Операция", "Сумма"], [10, 44, 16])
    for row, (platform, operation, amount) in enumerate(pnl.general, start=2):
        general.cell(row, 1, platform)
        general.cell(row, 2, operation)
        general.cell(row, 3, excel_value(amount, "money")).number_format = FORMATS["money"]

    problems = workbook.create_sheet("Проблемы")
    _header(problems, 1, ["Что проверить"], [140])
    for row, problem in enumerate(pnl.problems or ("Проблем не найдено",), start=2):
        problems.cell(row, 1, problem).alignment = Alignment(wrap_text=True, vertical="top")
    return workbook


def export_report(path: Path, pnl: Pnl) -> None:
    workbook = build_workbook(pnl)
    temporary = path.with_name(path.name + ".tmp")
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(temporary)
        os.replace(temporary, path)
    except PermissionError as exc:
        temporary.unlink(missing_ok=True)
        raise ExportError(f"Не удалось сохранить {path.name}: файл открыт в Excel? Закройте его и повторите.") from exc
    except OSError as exc:
        temporary.unlink(missing_ok=True)
        raise ExportError(f"Не удалось сохранить {path.name}: {exc}") from exc
